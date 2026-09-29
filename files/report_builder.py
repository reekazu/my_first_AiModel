# -*- coding: utf-8 -*-
"""Формирует итоговый Excel-файл сверки с раскраской расхождений."""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from reconcile_core import ReconciliationResult

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(color="FFFFFF", bold=True)
MISMATCH_FILL = PatternFill("solid", fgColor="FFC7CE")
ONLY_ONE_SIDE_FILL = PatternFill("solid", fgColor="FFEB9C")
OK_FILL = PatternFill("solid", fgColor="C6EFCE")


def _write_df(ws, df: pd.DataFrame, start_row: int = 1, highlight: str | None = None):
    if df.empty:
        ws.cell(row=start_row, column=1, value="Нет данных").font = Font(italic=True)
        return start_row + 1

    for c, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=start_row, column=c, value=str(col_name))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    fill = {
        "mismatch": MISMATCH_FILL,
        "only": ONLY_ONE_SIDE_FILL,
        "ok": OK_FILL,
    }.get(highlight)

    for r, (_, row) in enumerate(df.iterrows(), start=start_row + 1):
        for c, val in enumerate(row.tolist(), start=1):
            cell = ws.cell(row=r, column=c, value=val)
            if fill:
                cell.fill = fill

    for c in range(1, len(df.columns) + 1):
        col_letter = get_column_letter(c)
        max_len = max(
            [len(str(df.columns[c - 1]))] + [len(str(v)) for v in df.iloc[:, c - 1].tolist()]
        )
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 40)

    return start_row + len(df) + 2


def build_report(result: ReconciliationResult, out_path: str, our_label: str, their_label: str):
    wb = Workbook()

    ws = wb.active
    ws.title = "Сводка"
    ws["A1"] = f"Сверка расчётов: {our_label} vs {their_label}"
    ws["A1"].font = Font(bold=True, size=14)
    row = 3
    for key, value in result.summary.items():
        ws.cell(row=row, column=1, value=key)
        ws.cell(row=row, column=2, value=value)
        row += 1
    diff = result.summary["Оборот по нашим данным (Дебет - Кредит)"] - result.summary[
        "Оборот по данным контрагента (Дебет - Кредит)"
    ]
    row += 1
    ws.cell(row=row, column=1, value="Итоговое расхождение оборота").font = Font(bold=True)
    cell = ws.cell(row=row, column=2, value=round(diff, 2))
    cell.font = Font(bold=True)
    if abs(diff) > 0.01:
        cell.fill = MISMATCH_FILL
    else:
        cell.fill = OK_FILL
    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["B"].width = 20

    ws2 = wb.create_sheet("Совпадения")
    _write_df(ws2, result.matched, highlight="ok")

    ws3 = wb.create_sheet("Отличается номер документа")
    _write_df(ws3, result.doc_mismatch, highlight="only")

    ws4 = wb.create_sheet("Расхождения по сумме или дате")
    _write_df(ws4, result.amount_mismatch, highlight="mismatch")

    ws5 = wb.create_sheet(f"Только у нас")
    _write_df(ws5, result.only_ours, highlight="only")

    ws6 = wb.create_sheet(f"Только у контрагента")
    _write_df(ws6, result.only_theirs, highlight="only")

    wb.save(out_path)
