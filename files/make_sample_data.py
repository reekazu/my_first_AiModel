# -*- coding: utf-8 -*-
"""Генерирует два тестовых файла, имитирующих выгрузку акта сверки из 1С,
с намеренными расхождениями — чтобы проверить работу reconcile_cli.py."""

from openpyxl import Workbook

OUR_FILE = "sample_our.xlsx"
THEIR_FILE = "sample_contractor.xlsx"

# (дата, документ, дебет, кредит)
our_rows = [
    ("01.02.2026", "Акт №1",  150000, 0),
    ("03.02.2026", "Акт №2",  220000, 0),
    ("05.02.2026", "Оплата №15", 0, 100000),
    ("10.02.2026", "Акт №3",  75000, 0),
    ("12.02.2026", "Оплата №16", 0, 220000),
    ("15.02.2026", "Акт №4",  60000, 0),   # будет отличаться суммой у контрагента
    ("20.02.2026", "Акт №5",  30000, 0),   # будет отсутствовать у контрагента
    ("22.02.2026", "Оплата №17", 0, 50000),
]

their_rows = [
    ("01.02.2026", "Акт 1",   150000, 0),      # номер оформлен по-другому, но совпадает
    ("03.02.2026", "Акт №2",  220000, 0),
    ("05.02.2026", "Оплата №15", 0, 100000),
    ("10.02.2026", "Акт №3",  75000, 0),
    ("12.02.2026", "Оплата №16", 0, 220000),
    ("15.02.2026", "Акт №4",  65000, 0),       # сумма отличается: 65000 вместо 60000
    ("22.02.2026", "Оплата №17", 0, 50000),
    ("25.02.2026", "Акт №6",  40000, 0),       # есть только у контрагента
]


def _write(path: str, org_name: str, rows: list[tuple]):
    wb = Workbook()
    ws = wb.active
    ws.title = "Акт сверки"
    ws["A1"] = f"Акт сверки взаимных расчётов — {org_name}"
    ws["A2"] = "Период: 01.02.2026 - 28.02.2026"
    header_row = 4
    headers = ["№ п/п", "Дата", "Документ", "Дебет", "Кредит"]
    for c, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=c, value=h)
    for i, (date, doc, debit, credit) in enumerate(rows, start=1):
        r = header_row + i
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=date)
        ws.cell(row=r, column=3, value=doc)
        ws.cell(row=r, column=4, value=debit if debit else None)
        ws.cell(row=r, column=5, value=credit if credit else None)
    total_row = header_row + len(rows) + 2
    ws.cell(row=total_row, column=3, value="Итого:")
    ws.cell(row=total_row, column=4, value=sum(r[2] for r in rows))
    ws.cell(row=total_row, column=5, value=sum(r[3] for r in rows))
    wb.save(path)


if __name__ == "__main__":
    _write(OUR_FILE, "ООО «Наша Компания»", our_rows)
    _write(THEIR_FILE, "ООО «Контрагент»", their_rows)
    print(f"Созданы файлы: {OUR_FILE}, {THEIR_FILE}")
